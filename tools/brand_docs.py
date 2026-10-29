# -*- coding: utf-8 -*-
"""
Skład dokumentów DOCX w identyfikacji „STOP DYSKRYMINACJI — logo i layout publikacji".

Wejście: pliki DOCX z pandoca (style Title, Subtitle, Heading 1–3, Body Text,
First Paragraph, Compact, Block Text, tabele). Wyjście: NOWY plik DOCX w katalogu
--out — oryginał zostaje nietknięty. TREŚĆ SIĘ NIE ZMIENIA: tytuł i podtytuł
przechodzą na okładkę, reszta zostaje w tej samej kolejności; skrypt sprawdza to
po zapisie (akapit po akapicie). Jedyna zmiana tekstu: etykieta spisu treści pandoca
„Table of Contents" → „Spis treści". Znaczniki redakcyjne i niezgodny numer wersji
skrypt tylko zgłasza („uwaga: …").

Co nadaje:
  - okładka jak w projekcie (strona 5): pole w kolorze serii, sygnet, nazwa, numer,
    nadtytuł, tytuł, podtytuł, dwie kreski; niżej odbiorcy, oryginalny zestaw znaków FE
    z tekstem alternatywnym z Załącznika nr 10 i zdanie o dofinansowaniu (jak w stopce
    serwisu). Bez tabel układu — w PDF okładka nie jest ogłaszana jako tabela;
  - artykuły (A01–A12): zestaw znaków i pas tytułowy na pierwszej stronie, bez okładki;
  - A4, marginesy, nagłówek „PORADNIK 02 · STOP DYSKRYMINACJI" i numer strony;
  - tekst 11,5 pt / interlinia 1,5, do lewej (projekt: Literata), nagłówki (projekt:
    Onest) — gdy fontów projektu nie ma w systemie: Georgia i Segoe UI;
  - ramki „Block Text" (noty, „Jak dostosować") na tle w kolorze serii; tabele
    z nagłówkiem w kolorze serii, powtarzanym na kolejnych stronach;
  - język dokumentu pl-PL i tytuł w metadanych — dla czytników ekranu i PDF z tagami.

Użycie:
  python tools/brand_docs.py PLIK.docx [...] --out KATALOG
  python tools/brand_docs.py --all KATALOG_ŹRÓDŁOWY --out KATALOG [--ochra] [--publikacja]

Bez --publikacja: wersje do przeglądu, nazwa jak źródło + „_identyfikacja". Z --publikacja:
nazwa bez oznaczenia wersji („P1_Kodeks_Rownosci_i_Inkluzywnosci.docx") — gotowa do wgrania
w Media (przypisanie po kodzie); dokument, w którym zostały znaczniki redakcyjne
([DO UZUPEŁNIENIA], [KONTROLA EKSPERTA]…) albo „wersja 0.1" w pliku v0-2, nie jest zapisywany.
PDF: tools/docx_to_pdf.ps1.

Paczki v2.0 (MD i gotowe PDF): wykrywane automatycznie, bez ponownego składu PDF.
  python tools/brand_docs.py --all KATALOG_Z_MD --out KATALOG --publikacja --szkice-cms
    --wtyczka KATALOG_STOP_DYSKRYMINACJI_CORE
W tym trybie --publikacja usuwa numer wersji WYŁĄCZNIE z nazw plików.
Status CMS to zawsze szkic; daty z kalendarza nie są terminami automatycznej publikacji.
"""
import argparse
import copy
import glob
import io
import os
import re
import sys

# v2 packages contain reviewed Markdown and already typeset PDFs, not DOCX.
# --publikacja only controls filenames; the CMS package always contains drafts.
if __name__ == '__main__':
    from pathlib import Path
    _task_args = sys.argv[1:]
    _task_all = _task_args[_task_args.index('--all') + 1] if '--all' in _task_args and _task_args.index('--all') + 1 < len(_task_args) else None
    _task_markdown = any(str(x).lower().endswith('.md') for x in _task_args)
    if _task_all and Path(_task_all).is_dir():
        _task_markdown = _task_markdown or any(Path(_task_all).rglob('[AP]*_*.md'))
    if '--szkice-cms' in _task_args or _task_markdown:
        from editorial_content import main as _task_editorial_main
        raise SystemExit(_task_editorial_main())

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_brand_icons as sign   # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BRAND = os.path.join(REPO, 'wp-content', 'plugins', 'stop-dyskryminacji-core', 'assets', 'brand')
if not os.path.isdir(BRAND):
    BRAND = os.path.join(REPO, 'stop-dyskryminacji', 'app', 'public', 'wp-content', 'plugins',
                         'stop-dyskryminacji-core', 'assets', 'brand')
LOCKUP = os.path.join(BRAND, 'zestawienie-znakow-pl-original.png')
LOCKUP_ALT = ('Znaki: Fundusze Europejskie dla Wielkopolski, barwy Rzeczypospolitej Polskiej, '
              'Dofinansowane przez Unię Europejską, Herb Województwa Wielkopolskiego')
FUNDING = ('Projekt „STOP DYSKRYMINACJI" nr FEWP.06.10-IP.01-0136/25 jest dofinansowany przez Unię '
           'Europejską ze środków Europejskiego Funduszu Społecznego Plus w ramach programu Fundusze '
           'Europejskie dla Wielkopolski 2021–2027.')
DECORATIVE_NS = 'http://schemas.microsoft.com/office/drawing/2017/decorative'
COVER_STYLE = 'SD Okładka'          # akapity dodane przez skład — pomijane w kontroli treści

INK, MUTED, BORDER, WHITE = '14181F', '4A5160', 'DCD6CB', 'FFFFFF'
MORELA, BAR_LIGHT = '#F59E6B', '#F1EEE8'

# Kolory serii — jak na stronie (klasy sd-seria-* w motywie). „soft" — tło ramek.
SERIES = {
    'P1': dict(num='01', color='0B5563', soft='E6F0F1', eyebrow='DOKUMENT WZORCOWY', label='Kodeks 01',
               audience=('Dla pracodawców', 'i działów HR')),
    'P2': dict(num='02', color='9A3412', soft='FCEBDF', eyebrow='PORADNIK', label='Poradnik 02',
               audience=('Dla osób pracujących', 'i kandydatów')),
    'P3': dict(num='03', color='9A3412', soft='FCEBDF', eyebrow='PORADNIK', label='Poradnik 03',
               audience=('Dla świadków', 'i współpracowników')),
    'P4': dict(num='04', color='2F6B3A', soft='E7F1EA', eyebrow='KOMPENDIUM', label='Kompendium 04',
               audience=('Dla pracodawców', 'i kadry zarządzającej')),
    'P0': dict(num='', color='0B5563', soft='E6F0F1', eyebrow='DOKUMENT ROBOCZY', label='Brief merytoryczny',
               audience=('Materiał zespołu projektu', '')),
    'P5': dict(num='', color='0B5563', soft='E6F0F1', eyebrow='DOKUMENT ROBOCZY', label='Strategia treści',
               audience=('Materiał zespołu projektu', '')),
    'A': dict(num='', color='0B5563', soft='E6F0F1', eyebrow='BAZA WIEDZY · ARTYKUŁ', label='Artykuł',
              audience=('', '')),
}
OCHRA = dict(color='8A5A00', soft='F6EEDC')       # do wyboru zamiast ceglanego dla P3 (--ochra)


def font_installed(name):
    dirs = [os.path.join(os.environ.get('WINDIR', 'C:/Windows'), 'Fonts'),
            os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Microsoft', 'Windows', 'Fonts')]
    key = name.lower().replace(' ', '')
    return any(key in f.lower().replace(' ', '') for d in dirs if os.path.isdir(d) for f in os.listdir(d))


HEAD = 'Onest' if font_installed('Onest') else 'Segoe UI'
BODY = 'Literata' if font_installed('Literata') else 'Georgia'


def lum(h):
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a, b):
    x, y = lum(a), lum(b)
    return (max(x, y) + 0.05) / (min(x, y) + 0.05)


# --------------------------------------------------------------------------- XML

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


# Kolejność elementów wg schematu OOXML — Word odrzuca plik („nieczytelna zawartość"),
# gdy np. w:shd stoi za w:ind albo w:tblCaption przed w:tblW.
ORDER = {
    'w:pPr': ['pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr', 'widowControl', 'numPr',
              'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens', 'kinsoku', 'wordWrap',
              'overflowPunct', 'topLinePunct', 'autoSpaceDE', 'autoSpaceDN', 'bidi', 'adjustRightInd',
              'snapToGrid', 'spacing', 'ind', 'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
              'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl', 'divId', 'cnfStyle', 'rPr',
              'sectPr', 'pPrChange'],
    'w:rPr': ['rStyle', 'rFonts', 'b', 'bCs', 'i', 'iCs', 'caps', 'smallCaps', 'strike', 'dstrike', 'outline',
              'shadow', 'emboss', 'imprint', 'noProof', 'snapToGrid', 'vanish', 'webHidden', 'color', 'spacing',
              'w', 'kern', 'position', 'sz', 'szCs', 'highlight', 'u', 'effect', 'bdr', 'shd', 'fitText',
              'vertAlign', 'rtl', 'cs', 'em', 'lang', 'eastAsianLayout', 'specVanish', 'oMath'],
    'w:tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
                'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar',
                'tblLook', 'tblCaption', 'tblDescription'],
    'w:trPr': ['cnfStyle', 'divId', 'gridBefore', 'gridAfter', 'wBefore', 'wAfter', 'cantSplit', 'trHeight',
               'tblHeader', 'tblCellSpacing', 'jc', 'hidden'],
    'w:tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
               'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
    'w:pBdr': ['top', 'left', 'bottom', 'right', 'between', 'bar'],
    'w:tblBorders': ['top', 'left', 'start', 'bottom', 'right', 'end', 'insideH', 'insideV'],
    'w:tblCellMar': ['top', 'left', 'start', 'bottom', 'right', 'end'],
}
ORDER_Q = {qn(k): [qn('w:' + x) for x in v] for k, v in ORDER.items()}


def child(parent, tag, **attrs):
    """Zwraca (albo tworzy we właściwym miejscu) element potomny i ustawia atrybuty."""
    c = parent.find(qn(tag))
    if c is None:
        c = el(tag)
        seq = ORDER_Q.get(parent.tag)
        if seq and qn(tag) in seq:
            later = set(seq[seq.index(qn(tag)) + 1:])
            nxt = next((x for x in parent if x.tag in later), None)
            if nxt is not None:
                nxt.addprevious(c)
            else:
                parent.append(c)
        else:
            parent.append(c)
    for k, v in attrs.items():
        c.set(qn(k), str(v))
    return c


def set_fonts(rpr, name):
    rf = child(rpr, 'w:rFonts')
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def set_color(rpr, hexcolor):
    c = child(rpr, 'w:color')
    for a in ('w:themeColor', 'w:themeShade', 'w:themeTint'):
        if c.get(qn(a)) is not None:
            del c.attrib[qn(a)]
    c.set(qn('w:val'), hexcolor)


def set_border(ppr, side, color, sz=8, space=4, val='single'):
    b = child(child(ppr, 'w:pBdr'), 'w:' + side)
    for k, v in (('w:val', val), ('w:sz', sz), ('w:space', space), ('w:color', color)):
        b.set(qn(k), str(v))


def shade(container, fill):
    s = child(container, 'w:shd')
    s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto')
    s.set(qn('w:fill'), fill)


# --------------------------------------------------------------------------- style

def restyle(doc, s, article):
    color = s['color']
    st = doc.styles
    # Pandoc zapisuje „Heading 1" wielką literą — wyszukiwanie po nazwie w python-docx tego nie znajduje.
    by_name = {x.name: x for x in st}
    names = set(by_name)

    dd = st.element.find(qn('w:docDefaults'))
    rpr_def = dd.find(qn('w:rPrDefault'))
    rpr = rpr_def.find(qn('w:rPr'))
    if rpr is None:
        rpr = el('w:rPr')
        rpr_def.append(rpr)
    set_fonts(rpr, BODY)
    lang = child(rpr, 'w:lang')
    lang.set(qn('w:val'), 'pl-PL')
    lang.set(qn('w:eastAsia'), 'pl-PL')
    lang.set(qn('w:bidi'), 'ar-SA')

    def para(name, font, size, bold=False, col=INK, before=0, after=6, line=1.5, keep=False):
        if name not in names:
            return None
        sty = by_name[name]
        sty.font.size = Pt(size)
        sty.font.bold = bold
        sty.font.italic = False
        set_fonts(sty.element.get_or_add_rPr(), font)
        set_color(sty.element.get_or_add_rPr(), col)
        pf = sty.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = line
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        if keep:
            pf.keep_with_next = True
        return sty

    para('Normal', BODY, 11.5, after=6)
    para('Body Text', BODY, 11.5, after=7)
    para('First Paragraph', BODY, 11.5, after=7)
    para('Compact', BODY, 11.5, after=4)

    h1 = para('Heading 1', HEAD, 30 if not article else 24, bold=True, col=color, after=14, line=1.1, keep=True)
    if h1 is not None and not article:
        h1.paragraph_format.page_break_before = True       # Kodeks: część właściwa od nowej strony
    h2 = para('Heading 2', HEAD, 18 if not article else 15, bold=True, col=color, before=26, after=8, line=1.15, keep=True)
    if h2 is not None:
        set_border(h2.element.get_or_add_pPr(), 'top', color, sz=12, space=10)
    para('Heading 3', HEAD, 13, bold=True, before=16, after=5, line=1.2, keep=True)
    para('Heading 4', HEAD, 11.5, bold=True, before=12, after=4, line=1.25, keep=True)

    # Spis treści (pole TOC z pandoca): nagłówek i pozycje krojem nagłówków. Pozycje Word
    # tworzy przy aktualizacji pola stylami wbudowanymi „toc 1"–„toc 3" (domyślnie tekst
    # szeryfowy) — ich definicja zapisana w pliku ma pierwszeństwo.
    para('TOC Heading', HEAD, 20, bold=True, col=color, before=0, after=12, line=1.15)
    for lvl, size, bold, before in ((1, 10.5, True, 6), (2, 10, False, 0), (3, 9.5, False, 0)):
        name = next((n for n in ('toc %d' % lvl, 'TOC %d' % lvl) if n in names), 'toc %d' % lvl)
        if name not in names:
            sty = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            sty.element.set(qn('w:styleId'), 'TOC%d' % lvl)
            sty.base_style = by_name['Normal']
            by_name[name] = sty
            names.add(name)
        para(name, HEAD, size, bold=bold, before=before, after=2, line=1.3)
        by_name[name].paragraph_format.left_indent = Cm(0.45 * (lvl - 1))

    bt = para('Block Text', HEAD, 10.5, before=8, after=8, line=1.4)
    if bt is not None:
        ppr = bt.element.get_or_add_pPr()
        for side in ('top', 'bottom', 'right'):
            set_border(ppr, side, s['soft'], sz=4, space=6)
        set_border(ppr, 'left', color, sz=24, space=8)
        shade(ppr, s['soft'])
        ind = child(ppr, 'w:ind')
        ind.set(qn('w:left'), str(int(Cm(0.35).twips)))
        ind.set(qn('w:right'), str(int(Cm(0.35).twips)))

    if COVER_STYLE not in names:
        cs = st.add_style(COVER_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        cs.base_style = by_name['Normal']
        cs.paragraph_format.space_before = Pt(0)
        cs.paragraph_format.space_after = Pt(0)
        cs.paragraph_format.line_spacing = 1.15


def style_tables(doc, s):
    for t in doc.tables:
        tblpr = t._tbl.tblPr
        w = child(tblpr, 'w:tblW')
        w.set(qn('w:type'), 'pct')
        w.set(qn('w:w'), '5000')
        borders = child(tblpr, 'w:tblBorders')
        for side, val in (('top', 'single'), ('bottom', 'single'), ('insideH', 'single'),
                          ('left', 'nil'), ('right', 'nil'), ('insideV', 'nil')):
            b = child(borders, 'w:' + side)
            b.set(qn('w:val'), val)
            if val != 'nil':
                b.set(qn('w:sz'), '4')
                b.set(qn('w:color'), BORDER)
        mar = child(tblpr, 'w:tblCellMar')
        for side, cm in (('top', 0.1), ('bottom', 0.1), ('left', 0.18), ('right', 0.18)):
            m = child(mar, 'w:' + side)
            m.set(qn('w:w'), str(int(Cm(cm).twips)))
            m.set(qn('w:type'), 'dxa')
        for ri, row in enumerate(t.rows):
            if ri == 0:
                child(row._tr.get_or_add_trPr(), 'w:tblHeader')     # powtarzany, w PDF jako nagłówek (TH)
            for cell in row.cells:
                if ri == 0:
                    shade(cell._tc.get_or_add_tcPr(), s['color'])
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.25
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        set_fonts(r._r.get_or_add_rPr(), HEAD)
                        if ri == 0:
                            r.font.bold = True
                            set_color(r._r.get_or_add_rPr(), WHITE)


# --------------------------------------------------------------------------- obrazy

def png_bytes(img):
    buf = io.BytesIO()
    img.save(buf, format='PNG', optimize=True)
    buf.seek(0)
    return buf


def bars_png(width_px=1100):
    from PIL import Image, ImageDraw
    ss, h = 2, int(width_px * 0.2)
    img = Image.new('RGBA', (width_px * ss, h * ss), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    bar_h = int(h * 0.36) * ss
    for i, c in enumerate((BAR_LIGHT, MORELA)):
        top = i * (h * ss - bar_h)
        d.rounded_rectangle([0, top, width_px * ss - 1, top + bar_h], radius=bar_h // 2, fill=c)
    return img.resize((width_px, h))


def mark_decorative(run):
    """Obraz ozdobny — w PDF z Worda 365 jako artefakt; czytnik ekranu go pomija."""
    for docpr in run._r.iter(qn('wp:docPr')):
        docpr.set('descr', '')
        ext_lst = etree.SubElement(docpr, qn('a:extLst'))
        ext = etree.SubElement(ext_lst, qn('a:ext'))
        ext.set('uri', '{C183D7F6-B498-43B3-948B-1728B52AA6E4}')
        dec = etree.SubElement(ext, '{%s}decorative' % DECORATIVE_NS, nsmap={'adec': DECORATIVE_NS})
        dec.set('val', '1')


def set_alt(run, text):
    for docpr in run._r.iter(qn('wp:docPr')):
        docpr.set('descr', text)


def text_run(p, text, size, color, bold=False, font=None, spacing=None):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    rpr = r._r.get_or_add_rPr()
    set_fonts(rpr, font or HEAD)
    set_color(rpr, color)
    if spacing:
        child(rpr, 'w:spacing').set(qn('w:val'), str(spacing))
    return r


def eyebrow_color(series_hex):
    for cand in ('F8C9A4', 'FBE3D1', 'FDEFE5', 'FFFFFF'):
        if contrast(cand, series_hex) >= 4.5:
            return cand
    return 'FFFFFF'


# --------------------------------------------------------------------------- okładka

def background(paragraph, hex_color, width, height):
    """Pole w kolorze serii: obraz ozdobny za tekstem, zakotwiczony w lewym górnym rogu strony
    (w PDF artefakt — czytnik ekranu go pomija)."""
    from PIL import Image
    img = Image.new('RGB', (int(width.cm * 40), int(height.cm * 40)), '#' + hex_color)
    run = paragraph.add_run()
    run.add_picture(png_bytes(img), width=width, height=height)
    mark_decorative(run)
    inline = run._r.find('.//' + qn('wp:inline'))
    anchor = OxmlElement('wp:anchor')
    for k, v in (('distT', 0), ('distB', 0), ('distL', 0), ('distR', 0), ('simplePos', 0),
                 ('relativeHeight', 251658240), ('behindDoc', 1), ('locked', 1), ('layoutInCell', 1),
                 ('allowOverlap', 1)):
        anchor.set(k, str(v))
    sp = OxmlElement('wp:simplePos')
    sp.set('x', '0')
    sp.set('y', '0')
    anchor.append(sp)
    for axis in ('positionH', 'positionV'):
        pos = OxmlElement('wp:' + axis)
        pos.set('relativeFrom', 'page')
        off = OxmlElement('wp:posOffset')
        off.text = '0'
        pos.append(off)
        anchor.append(pos)
    anchor.append(inline.find(qn('wp:extent')))
    anchor.append(OxmlElement('wp:wrapNone'))
    for tag in ('wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic'):
        x = inline.find(qn(tag))
        if x is not None:
            anchor.append(x)
    inline.getparent().replace(inline, anchor)


COVER_FIELD_H = Cm(19.5)


def cover_height_cm(title, subtitle, text_w):
    """Szacunkowa wysokość tekstu w polu okładki (metryki Segoe UI z zapasem 5%) — do ostrzeżenia,
    gdyby po zmianach tytuł lub podtytuł nie mieścił się w polu. None, gdy brak fontu."""
    try:
        from PIL import ImageFont
        fdir = os.path.join(os.environ.get('WINDIR', 'C:/Windows'), 'Fonts')
        bold = ImageFont.truetype(os.path.join(fdir, 'segoeuib.ttf'), 300)
        reg = ImageFont.truetype(os.path.join(fdir, 'segoeui.ttf'), 300)
    except OSError:
        return None

    def lines(text, font, size):
        max_w = text_w / 12700 * 300 / size / 1.05          # EMU → pt → px przy 300 px/pt
        n, cur = 1, ''
        for word in text.split():
            cand = (cur + ' ' + word).strip()
            if cur and font.getlength(cand) > max_w:
                n, cur = n + 1, word
            else:
                cur = cand
        return n

    asc, desc = bold.getmetrics()
    em = (asc + desc) / 300                       # wysokość wiersza „pojedynczo" w em
    pt = 42.5 + 0.62 / 2.54 * 72 + 70 + 9.5 * em * 1.15 + 8
    pt += lines(title, bold, 30) * 30 * em * 1.08 + 10
    if subtitle:
        pt += lines(subtitle, reg, 11.5) * 11.5 * em * 1.3
    pt += 34 + 9.6 * 0.2 / 2.54 * 72
    return pt / 72 * 2.54


def build_cover(doc, s, title, subtitle, code, article, h1_el=None):
    """Elementy okładki (albo pasa tytułowego artykułu) — tworzone na końcu dokumentu,
    zwracane do przeniesienia na początek. Bez tabel: tabela użyta do układu byłaby w PDF
    oznaczona jako tabela danych i czytnik ekranu ogłaszałby okładkę jako „tabelę 1×1"."""
    color = s['color']
    full_w = Cm(21.0) if not article else Cm(16.0)
    pad = Cm(1.8) if not article else Cm(0.9)
    els = []

    def cover_par():
        p = doc.add_paragraph()
        p.style = doc.styles[COVER_STYLE]
        els.append(p._p)
        return p

    def sign_line(p):
        run = p.add_run()
        run.add_picture(png_bytes(sign.png(220, '#FFFFFF', '#' + color, MORELA)), height=Cm(0.62))
        mark_decorative(run)
        text_run(p, '  STOP DYSKRYMINACJI', 9.5, WHITE, bold=True, spacing=8)
        if s['num']:
            text_run(p, '\t' + s['num'], 10, WHITE, bold=True)

    def eyebrow(p):
        text_run(p, s['eyebrow'] + ((' ' + code) if code else ''), 9.5, eyebrow_color(color), bold=True, spacing=30)

    if article:
        lp = cover_par()
        lp.paragraph_format.space_after = Pt(10)
        run = lp.add_run()
        run.add_picture(LOCKUP, width=full_w)
        set_alt(run, LOCKUP_ALT)

        # Pas tytułowy: akapity z jednakowym tłem, obramowaniem i wcięciami — Word składa je
        # w jedno pole. H1 zostaje zwykłym nagłówkiem (w PDF znacznik H1).
        def band(p_el, before=0, after=0):
            ppr = p_el.get_or_add_pPr()
            for x in ppr.findall(qn('w:pageBreakBefore')):
                ppr.remove(x)
            child(ppr, 'w:keepNext')
            for side, space in (('top', 16), ('left', 25), ('bottom', 16), ('right', 25)):
                set_border(ppr, side, color, sz=4, space=space)
            shade(ppr, color)
            sp_el = child(ppr, 'w:spacing')
            sp_el.set(qn('w:before'), str(int(Pt(before).twips)))
            sp_el.set(qn('w:after'), str(int(Pt(after).twips)))
            ind = child(ppr, 'w:ind')
            ind.set(qn('w:left'), str(int(pad.twips)))
            ind.set(qn('w:right'), str(int(pad.twips)))

        p = cover_par()
        sign_line(p)
        band(p._p)
        eye = cover_par()
        eyebrow(eye)
        band(eye._p, before=16, after=4)
        if h1_el is not None:
            for r in h1_el.iter(qn('w:r')):
                rpr = r.find(qn('w:rPr'))
                if rpr is None:
                    rpr = el('w:rPr')
                    r.insert(0, rpr)
                set_color(rpr, WHITE)
            band(h1_el)
            els.append(h1_el)

        fp = cover_par()
        fp.paragraph_format.space_before = Pt(4)
        fp.paragraph_format.space_after = Pt(14)
        text_run(fp, FUNDING, 7.5, MUTED)
        return els

    # Okładka publikacji: tekst pola w ramce o stałej wysokości (19,5 cm od góry strony),
    # kolor pola — obraz ozdobny za tekstem; niżej odbiorcy, zestaw znaków i zdanie o dofinansowaniu.
    frame = (('w:w', int(full_w.twips)), ('w:h', int(COVER_FIELD_H.twips)), ('w:hRule', 'exact'),
             ('w:wrap', 'notBeside'), ('w:hAnchor', 'page'), ('w:vAnchor', 'page'), ('w:x', 0), ('w:y', 0))

    def framed(before=0, after=0, line=1.15):
        p = cover_par()
        fr = child(p._p.get_or_add_pPr(), 'w:framePr')
        for k, v in frame:
            fr.set(qn(k), str(v))
        pf = p.paragraph_format
        pf.left_indent = pf.right_indent = pad
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
        return p

    # Górny margines pola: pusty akapit o stałej wysokości (odstęp „przed" pierwszego akapitu
    # ramki Word odkłada nad ramką, poza stronę — i skraca o niego pole).
    spacer = framed(line=Pt(42.5))
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    p = framed()
    p.paragraph_format.tab_stops.add_tab_stop(full_w - pad, WD_TAB_ALIGNMENT.RIGHT)
    sign_line(p)
    eyebrow(framed(before=70, after=8))
    text_run(framed(after=10, line=1.08), title, 30, WHITE, bold=True)
    if subtitle:
        text_run(framed(line=1.3), subtitle, 11.5, WHITE)
    run = framed(before=34).add_run()
    run.add_picture(png_bytes(bars_png()), width=Cm(9.6))
    mark_decorative(run)

    ap = cover_par()
    background(ap, color, full_w, COVER_FIELD_H)
    ap.paragraph_format.left_indent = ap.paragraph_format.right_indent = pad
    ap.paragraph_format.space_before = Pt(16)
    ap.paragraph_format.line_spacing = 1.3
    ap.paragraph_format.tab_stops.add_tab_stop(full_w - pad, WD_TAB_ALIGNMENT.RIGHT)
    a1, a2 = s['audience']
    text_run(ap, a1, 9.5, INK)
    text_run(ap, '\tPublikacja bezpłatna', 9.5, MUTED)
    ap.add_run().add_break()
    text_run(ap, a2, 9.5, INK)
    text_run(ap, '\tstopdyskryminacji.pl', 9.5, MUTED)

    lp = cover_par()
    lp.paragraph_format.left_indent = pad
    lp.paragraph_format.space_before = Pt(118)
    lp.paragraph_format.space_after = Pt(4)
    run = lp.add_run()
    run.add_picture(LOCKUP, width=full_w - 2 * pad)
    set_alt(run, LOCKUP_ALT)

    fp = cover_par()
    fp.paragraph_format.left_indent = fp.paragraph_format.right_indent = pad
    fp.paragraph_format.line_spacing = 1.25
    text_run(fp, FUNDING, 8, MUTED)
    return els


def page_field(p):
    r = p.add_run()
    for kind, txt in (('begin', None), (None, ' PAGE '), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            r._r.append(el('w:fldChar', **{'w:fldCharType': kind}))
        elif txt.strip() == 'PAGE':
            it = el('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            r._r.append(it)
        else:
            t = el('w:t')
            t.text = txt
            r._r.append(t)
    r.font.size = Pt(8.5)
    set_fonts(r._r.get_or_add_rPr(), HEAD)
    set_color(r._r.get_or_add_rPr(), INK)


def header_footer(sec, s, label, short):
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    text_run(hp, label.upper() + ' · STOP DYSKRYMINACJI', 7.5, s['color'], bold=True, spacing=16)
    text_run(hp, '\t' + short, 7.5, MUTED)
    set_border(hp._p.get_or_add_pPr(), 'bottom', BORDER, sz=4, space=6)
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
    text_run(fp, 'stopdyskryminacji.pl', 7.5, MUTED)
    fp.add_run('\t')
    page_field(fp)


# --------------------------------------------------------------------------- kontrola treści

def para_texts(root, skip_style_id=None, skip=()):
    out = []
    for p in root.iter(qn('w:p')):
        if p in skip:
            continue
        if skip_style_id:
            ps = p.find(qn('w:pPr') + '/' + qn('w:pStyle'))
            if ps is not None and ps.get(qn('w:val')) == skip_style_id:
                continue
        t = ''.join(x.text or '' for x in p.iter(qn('w:t')))
        if t.strip():
            out.append(t)
    return out


# Znaczniki redakcyjne wersji roboczych — przed publikacją powinny zniknąć. [DOSTOSUJ] w Kodeksie
# jest częścią wzoru (miejsce do uzupełnienia przez pracodawcę), dlatego liczony jest osobno.
MARKER = re.compile(r'\[([^\[\]\n]{0,120}?)\]')
EDITORIAL = re.compile(r'UZUPEŁNIENIA|WERYFIKACJI|KONTROLA EKSPERTA|DATA AUTORYZACJI|DECYZJA|ZESTAW ZNAKÓW'
                       r'|do potwierdzenia|publikator', re.I)


def draft_warnings(src, texts, doc, subtitle=""):
    """Uwagi do treści (skrypt jej nie zmienia): znaczniki redakcyjne i numer wersji z podtytułu
    albo z wiersza „Wersja" tabeli metryczki, niezgodny z nazwą pliku. Lista par
    (uwaga, czy wstrzymuje tryb --publikacja)."""
    out = []
    marks = [m.group(1) for t in texts for m in MARKER.finditer(t)]
    ed = [x for x in marks if EDITORIAL.search(x)]
    adj = [x for x in marks if x.upper().startswith('DOSTOSUJ')]
    if ed:
        kinds = sorted({re.split(r' — |:', x)[0].strip() for x in ed})
        out.append(('%d znaczników redakcyjnych do usunięcia przed publikacją (%s)' % (len(ed), ', '.join('[%s]' % k for k in kinds[:4]) + (' …' if len(kinds) > 4 else '')), True))
    if adj:
        out.append(('%d miejsc [DOSTOSUJ]' % len(adj), False))
    m = re.search(r'_v(\d+)-(\d+)', os.path.basename(src))
    if m:
        fv = '%s.%s' % m.groups()
        found = set(re.findall(r'[Ww]ersj\w*\s+(\d+[.,]\d+)', subtitle))
        for t in doc.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                mm = len(cells) > 1 and cells[0].lower() == 'wersja' and re.match(r'(\d+[.,]\d+)', cells[1])
                if mm:
                    found.add(mm.group(1))
        bad = sorted({x.replace(',', '.') for x in found} - {fv})
        if bad:
            out.append(('w treści „wersja %s", a plik to v%s' % (', '.join(bad), fv.replace('.', '-')), True))
    return out


def publication_name(src):
    """Nazwa do wgrania na stronę: bez oznaczenia wersji roboczej („_v0-2"). Kod na początku
    (P1_, A01_) zostaje — po nim wtyczka przypisuje plik do publikacji."""
    base = os.path.splitext(os.path.basename(src))[0]
    base = re.sub(r'[_ -]+v\d+(?:[-.]\d+)*(?=$|[_ -])', '', base, flags=re.I)
    return re.sub(r'(?:[_ -]+(?:final\w*|ostateczn\w*|robocz\w*|draft|autoryzowan\w*|po[_ -]autoryzacji))+$', '', base, flags=re.I)


def kind_of(path):
    m = re.match(r'(P[0-9]|A\d\d)_', os.path.basename(path))
    if not m:
        return None, None
    k = m.group(1)
    return ('A', k) if k.startswith('A') else (k, k)


def short_title(title, n=60):
    t = re.split(r' — |: |\? ', title)[0]
    return t if len(t) <= n else t[:n - 1].rstrip() + '…'


def brand(src, out_dir, ochra=False, publish=False):
    kind, code = kind_of(src)
    if not kind:
        return None, [], 'pominięto (nie rozpoznano rodzaju z nazwy pliku)'
    s = dict(SERIES[kind])
    if ochra and kind == 'P3':
        s.update(OCHRA)
    article = kind == 'A'

    doc = Document(src)
    body = doc.element.body
    title_p = next((p for p in doc.paragraphs if p.style.name == 'Title'), None)
    sub_p = next((p for p in doc.paragraphs if p.style.name == 'Subtitle'), None)
    h1_p = next((p for p in doc.paragraphs if p.style.name == 'Heading 1'), None)
    title = title_p.text.strip() if title_p is not None else os.path.basename(src)
    subtitle = sub_p.text.strip() if sub_p is not None else ''
    removed = {e._p for e in (title_p, sub_p) if e is not None}
    expected = para_texts(body, skip=removed)          # treść bez tytułu i podtytułu (idą na okładkę)
    items = draft_warnings(src, [title, subtitle] + expected, doc, subtitle)
    warns = [w for w, _ in items]
    if publish and any(b for _, b in items):
        return None, warns, 'WSTRZYMANO — wersja robocza, nie zapisano pliku do publikacji'
    if not article:
        est = cover_height_cm(title, subtitle, Cm(21.0) - 2 * Cm(1.8))
        if est and est > COVER_FIELD_H.cm - 0.5:
            warns.append('tytuł z podtytułem może nie zmieścić się w polu okładki (ok. %.1f z %.1f cm) — sprawdź PDF'
                         % (est, COVER_FIELD_H.cm))

    # Jedyna zmiana tekstu: angielska etykieta spisu treści, wstawiana przez pandoca.
    notes = []
    toc_ids = {x.style_id for x in doc.styles if x.name == 'TOC Heading'}
    for p in body.iter(qn('w:p')):
        ps = p.find(qn('w:pPr') + '/' + qn('w:pStyle'))
        texts = list(p.iter(qn('w:t')))
        if ps is not None and ps.get(qn('w:val')) in toc_ids and ''.join(t.text or '' for t in texts) == 'Table of Contents':
            texts[0].text = 'Spis treści'
            for t in texts[1:]:
                t.text = ''
            expected = ['Spis treści' if x == 'Table of Contents' else x for x in expected]
            notes.append('„Table of Contents" → „Spis treści"')

    restyle(doc, s, article)
    style_tables(doc, s)
    for e in removed:
        body.remove(e)

    sec = doc.sections[-1]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2.4)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.1)

    h1_el = None
    if article and h1_p is not None:
        h1_el = h1_p._p
        body.remove(h1_el)
    cover = build_cover(doc, s, title, subtitle, code if article else '', article, h1_el)
    for e in cover:
        if e.getparent() is body:                      # H1 artykułu jest już wyjęty z treści
            body.remove(e)
    for i, e in enumerate(cover):
        body.insert(i, e)

    if not article:
        # Sekcja okładki: bez marginesów i bez nagłówka; treść od nowej strony.
        sect = copy.deepcopy(doc.sections[-1]._sectPr)
        for ref in sect.findall(qn('w:headerReference')) + sect.findall(qn('w:footerReference')):
            sect.remove(ref)
        pgmar = sect.find(qn('w:pgMar'))
        for a in ('w:top', 'w:left', 'w:right', 'w:bottom', 'w:header', 'w:footer', 'w:gutter'):
            pgmar.set(qn(a), '0')
        typ = sect.find(qn('w:type'))
        if typ is None:
            typ = el('w:type')
            sect.insert(0, typ)
        typ.set(qn('w:val'), 'nextPage')
        sp = el('w:p')
        ppr = el('w:pPr')
        ppr.append(el('w:pStyle', **{'w:val': doc.styles[COVER_STYLE].style_id}))
        ppr.append(sect)
        sp.append(ppr)
        body.insert(len(cover), sp)
        header_footer(doc.sections[-1], s, s['label'], short_title(title))
    else:
        doc.sections[-1].different_first_page_header_footer = True
        header_footer(doc.sections[-1], s, s['label'] + ' ' + code, short_title(title))

    cp = doc.core_properties
    cp.title = title
    cp.language = 'pl-PL'
    cp.subject = 'STOP DYSKRYMINACJI — FEWP.06.10-IP.01-0136/25'

    os.makedirs(out_dir, exist_ok=True)
    name = publication_name(src) if publish else os.path.splitext(os.path.basename(src))[0] + '_identyfikacja'
    out = os.path.join(out_dir, name + '.docx')
    doc.save(out)

    # Kontrola po zapisie: treść akapit po akapicie, tytuł i podtytuł na okładce.
    chk = Document(out)
    sid = chk.styles[COVER_STYLE].style_id
    after = para_texts(chk.element.body, skip_style_id=sid)
    cover_txt = ' '.join(para_texts(chk.element.body))
    if article:
        # W artykule pole Title („A1. …") powtarza H1; H1 stoi w pasie, numer — w nadtytule.
        on_cover = h1_el is not None and ''.join(x.text or '' for x in h1_el.iter(qn('w:t'))) in cover_txt and code in cover_txt
    else:
        on_cover = title in cover_txt and (not subtitle or subtitle in cover_txt)
    ok = after == expected and on_cover
    if ok:
        return out, warns, 'treść zgodna (%d akapitów), %s%s' % (len(after), 'tytuł (H1) i numer w pasie tytułowym' if article else 'tytuł%s na okładce' % (' i podtytuł' if subtitle else ''),
                                                         ('; ' + ', '.join(notes)) if notes else '')
    if after == expected:
        return out, warns, 'RÓŻNICA: tytuł nie trafił na okładkę'
    diff = next((i for i, (a, b) in enumerate(zip(after, expected)) if a != b), min(len(after), len(expected)))
    return out, warns, 'RÓŻNICA W TREŚCI od akapitu %d: %r vs %r' % (diff, (after[diff:diff + 1] or [''])[0][:60], (expected[diff:diff + 1] or [''])[0][:60])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='*')
    ap.add_argument('--all', help='katalog źródłowy: wszystkie P*_*.docx i A*_*.docx')
    ap.add_argument('--out', required=True)
    ap.add_argument('--ochra', action='store_true', help='P3 w ochrze zamiast ceglanego')
    ap.add_argument('--publikacja', action='store_true',
                    help='pliki do wgrania na stronę: nazwa bez „_v0-2"; dokument ze znacznikami redakcyjnymi '
                         'albo niezgodnym numerem wersji jest wstrzymywany')
    a = ap.parse_args()
    files = list(a.files)
    if a.all:
        files += sorted(glob.glob(os.path.join(a.all, 'P[0-9]_*.docx')) + glob.glob(os.path.join(a.all, 'A[0-9][0-9]_*.docx')))
    if not files:
        ap.error('Nie znaleziono dokumentów wejściowych; niczego nie zapisano.')
    print('Fonty: nagłówki %s, tekst %s%s' % (HEAD, BODY, '' if (HEAD, BODY) == ('Onest', 'Literata') else ' (fonty projektu Onest i Literata nie są zainstalowane)'))
    bad = 0
    for f in files:
        out, warns, msg = brand(f, a.out, ochra=a.ochra, publish=a.publikacja)
        bad += 'RÓŻNICA' in msg or 'WSTRZYMANO' in msg
        print('  %-58s %s' % (os.path.basename(f)[:58], msg))
        for w in warns:
            print('  %-58s   uwaga: %s' % ('', w))
    sys.exit(1 if bad else 0)


if __name__ == '__main__':
    main()
